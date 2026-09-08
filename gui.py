"""
Grafické rozhraní (Tkinter GUI) pro kontrolu a přejmenování výkresů.
Obsahuje: volbu CPU/GPU, volbu výkonu hardware, progress s %, scroll kolečkem a undo zálohu.
"""
import os
import json
import time
import queue
import re
import threading
import logging
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor
from PIL import Image, ImageTk, ImageOps, ImageEnhance

from config import load_config, save_config, get_dynamic_hardware_profiles, get_detected_devices
from ocr_engine import load_models, load_and_resize_image, process_image, prepare_crops

logger = logging.getLogger(__name__)


CHECKPOINT_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "session_checkpoint.json")
_MAX_PREVIEW_CACHE_SIZE = 30  # Maximální počet náhledů v paměti (LRU eviction)
_WINDOWS_INVALID_CHARS = re.compile(r'[<>:"/\\|?*]')  # Znaky zakázané ve Windows názvech souborů


def _get_undo_file(folder: str) -> str:
    """Per-folder undo soubor — why: původně jeden globální undo_backup.json přepisoval undo při paralelních jobech."""
    import hashlib
    h = hashlib.md5(os.path.abspath(folder).lower().encode()).hexdigest()[:8]
    base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, f"undo_backup_{h}.json")


def load_all_checkpoints() -> dict:
    """Načte všechny uložené kontrolní body — why: původně truncated open("w") + výpadek = 0B corrupt → return {} ztráta všech složek."""
    if not os.path.exists(CHECKPOINT_FILE):
        # zkusit zálohu pokud existuje
        bak = CHECKPOINT_FILE + ".bak"
        if os.path.exists(bak):
            try:
                with open(bak, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass
        return {}
    try:
        with open(CHECKPOINT_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        logger.warning("Nelze načíst session_checkpoint.json: %s — zkouším zálohu", e)
        # pokus o obnovu ze zálohy (.tmp je neúplný zápis – nečíst)
        for alt in (CHECKPOINT_FILE + ".bak",):
            if os.path.exists(alt):
                try:
                    with open(alt, "r", encoding="utf-8") as f:
                        data = json.load(f)
                    logger.info("Obnoveno ze zálohy %s", alt)
                    return data
                except Exception:
                    continue
        return {}


def save_all_checkpoints(data: dict):
    """Bezpečně uloží kontrolní body — why: původně open("w") truncated → výpadek = corrupt 0B."""
    try:
        tmp = CHECKPOINT_FILE + ".tmp"
        bak = CHECKPOINT_FILE + ".bak"
        # zálohovat předchozí verzi pokud existuje
        if os.path.exists(CHECKPOINT_FILE):
            try:
                import shutil
                shutil.copy2(CHECKPOINT_FILE, bak)
            except Exception:
                pass
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
            f.flush()
            try:
                os.fsync(f.fileno())
            except Exception:
                pass
        os.replace(tmp, CHECKPOINT_FILE)
    except Exception as e:
        logger.error("Nelze zapsat do session_checkpoint.json: %s", e)


def get_folder_checkpoint(folder_path: str) -> dict | None:
    """Vrátí uložený kontrolní bod pro danou složku, pokud existuje."""
    norm_path = os.path.abspath(folder_path).lower()
    all_chk = load_all_checkpoints()
    for stored_path, data in all_chk.items():
        if os.path.abspath(stored_path).lower() == norm_path:
            return data
    return None


def save_folder_checkpoint(folder_path: str, data: dict):
    """Uloží nebo aktualizuje kontrolní bod pro konkrétní složku."""
    norm_path = os.path.abspath(folder_path)
    all_chk = load_all_checkpoints()
    key_to_use = norm_path
    for stored_path in list(all_chk.keys()):
        if os.path.abspath(stored_path).lower() == norm_path.lower():
            key_to_use = stored_path
            break
    all_chk[key_to_use] = data
    save_all_checkpoints(all_chk)


def remove_folder_checkpoint(folder_path: str):
    """Odstraní záznam kontrolního bodu pro konkrétní složku."""
    norm_path = os.path.abspath(folder_path).lower()
    all_chk = load_all_checkpoints()
    keys_to_del = [k for k in all_chk if os.path.abspath(k).lower() == norm_path]
    if keys_to_del:
        for k in keys_to_del:
            del all_chk[k]
        save_all_checkpoints(all_chk)



def sanitize_filename(name: str) -> str:
    """Odstraní znaky zakázané ve Windows názvech souborů (<>:\"/\\|?*) a ořízne mezery."""
    if not name:
        return name
    cleaned = _WINDOWS_INVALID_CHARS.sub('', name)
    # Odstranění koncových teček a mezer (Windows je ignoruje v názvech)
    cleaned = cleaned.rstrip('. ')
    return cleaned.strip()


def clean_only_drawing_number(val: str) -> str:
    """Z textu názvu nebo souboru odstraní příponu i formát v závorce a vrátí pouze čisté číslo výkresu."""
    if not val:
        return ""
    # 1. Odstranění přípony souboru (.tif, .jpg, .png atd.)
    base = re.sub(r'\.(tiff?|jpg|jpeg|png)$', '', str(val).strip(), flags=re.IGNORECASE).strip()
    # 2. Odstranění formátu v závorce na konci (např. ' (A4)', ' (A5)', ' (A3+)')
    base = re.sub(r'\s*\([a-zA-Z0-9\+\-\s]+\)\s*$', '', base).strip()
    return base


def export_to_excel(arg1, arg2, paper_format: str = "A4") -> bool:
    """
    Vyexportuje výkresy do formátu Excel (.xlsx) s přesnými sloupci:
    Sloupec A: Číslo výkresu
    Sloupec B: Název výkresu
    Sloupec C: Formát
    """
    if isinstance(arg1, str):
        excel_path = arg1
        items = arg2
    else:
        items = arg1
        excel_path = str(arg2)

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if not items:
        logger.warning("Žádné položky k exportu do Excelu.")
        return False

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Výkresy"

    # Hlavička přesně dle specifikace: Číslo výkresu, Název výkresu, Formát
    headers = ["Číslo výkresu", "Název výkresu", "Formát"]
    for col_idx, h_text in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h_text)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0")
    )

    for col_idx in range(1, 4):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    data_font = Font(name="Segoe UI", size=10)
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_center = Alignment(horizontal="center", vertical="center")

    for row_idx, item in enumerate(items, start=2):
        if isinstance(item, dict):
            num_val = item.get("drawing_num") or item.get("number") or item.get("user_num") or item.get("num") or ""
            name_val = item.get("drawing_name") or item.get("name") or item.get("user_name") or item.get("user_title") or item.get("title") or ""
            fmt_val = item.get("drawing_format") or item.get("format") or item.get("paper_format") or paper_format
        elif isinstance(item, (list, tuple)):
            num_val = item[0] if len(item) > 0 else ""
            name_val = item[1] if len(item) > 1 else ""
            fmt_val = item[2] if len(item) > 2 else paper_format
        else:
            num_val, name_val, fmt_val = str(item), "", paper_format

        ws.cell(row=row_idx, column=1, value=str(num_val)).alignment = data_align_center
        ws.cell(row=row_idx, column=2, value=str(name_val)).alignment = data_align_left
        ws.cell(row=row_idx, column=3, value=str(fmt_val or paper_format)).alignment = data_align_center

        for col_idx in range(1, 4):
            ws.cell(row=row_idx, column=col_idx).font = data_font
            ws.cell(row=row_idx, column=col_idx).border = thin_border

    # Automatické přizpůsobení šířky sloupců
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 6, 18)

    ws.row_dimensions[1].height = 25
    for r in range(2, len(items) + 2):
        ws.row_dimensions[r].height = 20

    try:
        wb.save(excel_path)
        logger.info("Úspěšně uložen Excel soubor: %s (%d položek)", excel_path, len(items))
        return True
    except PermissionError:
        base, ext = os.path.splitext(excel_path)
        fallback_path = f"{base}_nove{ext}"
        try:
            wb.save(fallback_path)
            logger.warning("Soubor %s je uzamčen v Excelu. Uloženo jako %s", excel_path, fallback_path)
            return True
        except Exception:
            timestamp_path = f"{base}_{int(time.time())}{ext}"
            try:
                wb.save(timestamp_path)
                logger.warning("Soubor %s i %s uzamčen. Uloženo jako %s", excel_path, fallback_path, timestamp_path)
                return True
            except Exception as e:
                logger.error("Nelze uložit Excel soubor: %s", e)
                return False


class RenameApp:
    """Hlavní grafické rozhraní a orchestrace zpracování výkresů."""

    def __init__(self, root: tk.Tk):
        self.root = root
        self.config = load_config()
        self.root.withdraw()
        self.root.wm_attributes("-topmost", True)

        # 1. Výběr složky s výkresy
        self.folder = filedialog.askdirectory(
            parent=self.root,
            title="Vyberte složku s výkresy (TIFF, JPG, PNG)"
        )
        if not self.folder:
            self.on_exit()
            return

        # 2. Načtení seznamu souborů — why: NAS smazán/výpadek → os.listdir crash bez dialogu
        try:
            self.files = sorted([
                f for f in os.listdir(self.folder)
                if f.lower().endswith(('.tiff', '.tif', '.jpg', '.jpeg', '.png'))
            ])
        except Exception as e:
            logger.error("Nelze číst složku %s: %s", self.folder, e)
            try:
                messagebox.showerror("Chyba", f"Nelze číst složku:\n{self.folder}\n\n{e}\n\nZkontrolujte síť / práva a zkuste znovu.")
            except Exception:
                pass
            self.on_exit()
            return
        if not self.files:
            messagebox.showinfo("Informace", "Ve zvolené složce nebyly nalezeny žádné výkresy.")
            self.on_exit()
            return

        # 3. Kontrola existujícího rozpracovaného stavu (Checkpoint)
        self.cached_results = {}
        self.active_checkpoint = get_folder_checkpoint(self.folder)

        if self.active_checkpoint and "results" in self.active_checkpoint:
            valid_cached = {
                f: self.active_checkpoint["results"][f]
                for f in self.files
                if f in self.active_checkpoint["results"]
            }
            if valid_cached:
                action = self._prompt_resume_dialog(len(self.files), len(valid_cached), self.active_checkpoint)
                if action == "resume":
                    self.cached_results = valid_cached
                    self.paper_format, self.hardware_profile, self.device_choice = self._show_setup_dialog(
                        default_format=self.active_checkpoint.get("paper_format"),
                        default_profile=self.active_checkpoint.get("hardware_profile"),
                        default_device=self.active_checkpoint.get("device")
                    )
                elif action == "review_now":
                    self.cached_results = valid_cached
                    self.paper_format = self.active_checkpoint.get("paper_format", "A4")
                    self.hardware_profile = self.active_checkpoint.get("hardware_profile", "turbo")
                    self.device_choice = self.active_checkpoint.get("device", "auto")
                    self.root.wm_attributes("-topmost", False)
                    # Sestavení results a přímý přechod ke kontrole
                    self.results = [(f, data.get("num"), data.get("conf")) for f, data in valid_cached.items()]
                    self.preview_executor = ThreadPoolExecutor(max_workers=2)
                    self.preview_cache = OrderedDict()
                    self._preview_request_id = 0
                    self._current_photo = None
                    self.root.deiconify()
                    self.show_review_gui()
                    return
                elif action == "restart":
                    remove_folder_checkpoint(self.folder)
                    self.active_checkpoint = None
                    self.cached_results = {}
                    self.paper_format, self.hardware_profile, self.device_choice = self._show_setup_dialog()
                else:
                    self.on_exit()
                    return
            else:
                self.paper_format, self.hardware_profile, self.device_choice = self._show_setup_dialog()
        else:
            self.paper_format, self.hardware_profile, self.device_choice = self._show_setup_dialog()

        self.root.wm_attributes("-topmost", False)

        if not self.paper_format or not self.hardware_profile or not self.device_choice:
            self.on_exit()
            return

        # Uložení nastavení pro příště
        self.config["paper_format"] = self.paper_format
        self.config["hardware_profile"] = self.hardware_profile
        self.config["device"] = self.device_choice
        save_config(self.config)

        logger.info("Vybrána složka: %s (%d souborů), profil: %s, zařízení: %s",
                    self.folder, len(self.files), self.hardware_profile, self.device_choice)

        # 4. Zobrazení loading okna
        self.root.deiconify()
        self.root.title("Přejmenování Výkresů")
        self.root.geometry("560x240")

        self.lbl_status = tk.Label(root, text="Inicializace AI a OCR modelů...", font=("Arial", 11), pady=5)
        self.lbl_status.pack()

        self.progress = ttk.Progressbar(root, orient="horizontal", length=460, mode="indeterminate")
        self.progress.pack(pady=6)
        self.progress.start(15)

        self.lbl_pct = tk.Label(root, text="", font=("Arial", 10, "bold"), fg="#1565C0")
        self.lbl_pct.pack()

        prof_label = get_dynamic_hardware_profiles().get(self.hardware_profile, {}).get("short_name", self.hardware_profile)
        dev_label = "GPU" if self.device_choice == "cuda" else ("CPU" if self.device_choice == "cpu" else "Auto")
        self.lbl_counter = tk.Label(
            root, text=f"0 / {len(self.files)} souborů | Režim: {prof_label} | {dev_label}",
            font=("Arial", 9), fg="#666"
        )
        self.lbl_counter.pack()

        # Tlačítko pro bezpečné pozastavení a ukončení
        self._pause_requested = threading.Event()
        self.btn_pause = tk.Button(
            root,
            text="⏸ Pozastavit a bezpečně ukončit",
            font=("Segoe UI", 9, "bold"),
            bg="#ECEFF1",
            fg="#C62828",
            command=self.request_pause,
            pady=3,
            padx=8
        )
        self.btn_pause.pack(pady=(6, 0))

        self.results = []
        self.queue = queue.Queue()
        self.preview_executor = ThreadPoolExecutor(max_workers=2)
        self.preview_cache = OrderedDict()
        self._preview_request_id = 0
        self._current_photo = None

        # Časování pro ETA
        self._processing_start_time = None

        # Spuštění výpočtu ve vlákně
        self.root.protocol("WM_DELETE_WINDOW", self.request_pause)
        threading.Thread(target=self.process_files_thread, daemon=True).start()
        self.root.after(100, self.check_queue)

    def request_pause(self):
        """Vyžádá bezpečné zastavení zpracování po aktuálním výkresu."""
        if hasattr(self, '_pause_requested'):
            self._pause_requested.set()
        if hasattr(self, 'btn_pause') and self.btn_pause.winfo_exists():
            self.btn_pause.config(text="⏳ Ukládám postup a ukončuji...", state="disabled")
        if hasattr(self, 'lbl_status') and self.lbl_status.winfo_exists():
            self.lbl_status.config(text="Pozastavuji... Ukládám postup po dokončení aktuálního souboru.")

    def _prompt_resume_dialog(self, total_files_count: int, cached_count: int, checkpoint: dict) -> str | None:
        """
        Zobrazí přehledný dialog při nalezení dřívějšího rozpracovaného zpracování.
        Vrací: 'resume' | 'review_now' | 'restart' | None
        """
        dialog = tk.Toplevel(self.root)
        dialog.title("Rozpracované zpracování výkresů")
        dialog.geometry("560x360")
        dialog.resizable(False, False)
        dialog.grab_set()
        dialog.attributes("-topmost", True)
        dialog.focus_force()

        header_frame = tk.Frame(dialog, bg="#1976D2", pady=12, padx=16)
        header_frame.pack(fill="x")

        tk.Label(
            header_frame,
            text="⏸ Nalezeno rozpracované zpracování",
            font=("Segoe UI", 13, "bold"),
            bg="#1976D2",
            fg="white"
        ).pack(anchor="w")

        info_frame = tk.Frame(dialog, padx=20, pady=14)
        info_frame.pack(fill="both", expand=True)

        folder_name = os.path.basename(self.folder.rstrip("\\/")) or self.folder
        updated_at = checkpoint.get("updated_at", "Neznámý čas")
        pct = int((cached_count / total_files_count) * 100) if total_files_count > 0 else 0
        remaining_count = max(0, total_files_count - cached_count)
        status_text = "Připraveno ke kontrole" if checkpoint.get("status") == "review" else "Rozpracováno v průběhu"

        details = [
            ("📁 Složka výkresů:", folder_name),
            ("📊 Celkem nalezeno:", f"{total_files_count} souborů"),
            ("✅ Již zpracováno:", f"{cached_count} výkresů ({pct} %)"),
            ("⏳ Zbývá dopočítat:", f"{remaining_count} výkresů"),
            ("🕒 Poslední uložení:", f"{updated_at}  ({status_text})"),
        ]

        for r, (label_text, val_text) in enumerate(details):
            tk.Label(info_frame, text=label_text, font=("Segoe UI", 9, "bold"), fg="#37474F").grid(row=r, column=0, sticky="w", pady=2)
            tk.Label(info_frame, text=val_text, font=("Segoe UI", 9), fg="#1565C0" if r == 2 else "#263238").grid(row=r, column=1, sticky="w", padx=10, pady=2)

        prompt_lbl = tk.Label(
            info_frame,
            text="Můžete pokračovat tam, kde jste skončili, nebo přejít přímo ke kontrole:",
            font=("Segoe UI", 9, "italic"),
            fg="#555"
        )
        prompt_lbl.grid(row=len(details), column=0, columnspan=2, sticky="w", pady=(10, 0))

        result = [None]

        def on_choose(action):
            result[0] = action
            dialog.destroy()

        btn_frame = tk.Frame(dialog, bg="#ECEFF1", pady=10, padx=16)
        btn_frame.pack(fill="x", side="bottom")

        tk.Button(
            btn_frame,
            text=f"▶ Pokračovat ({remaining_count} zbývá)",
            bg="#2E7D32",
            fg="white",
            font=("Segoe UI", 9, "bold"),
            command=lambda: on_choose("resume"),
            padx=10,
            pady=4
        ).pack(side="left", padx=4)

        tk.Button(
            btn_frame,
            text=f"🔎 Zkontrolovat hotové ({cached_count})",
            bg="#1976D2",
            fg="white",
            font=("Segoe UI", 9, "bold"),
            command=lambda: on_choose("review_now"),
            padx=10,
            pady=4
        ).pack(side="left", padx=4)

        tk.Button(
            btn_frame,
            text="🔄 Od začátku",
            font=("Segoe UI", 9),
            command=lambda: on_choose("restart"),
            padx=6,
            pady=4
        ).pack(side="left", padx=4)

        tk.Button(
            btn_frame,
            text="Zrušit",
            font=("Segoe UI", 9),
            command=lambda: dialog.destroy(),
            padx=6,
            pady=4
        ).pack(side="right", padx=4)

        dialog.protocol("WM_DELETE_WINDOW", lambda: dialog.destroy())
        dialog.wait_window()
        return result[0]

    def _show_setup_dialog(self, default_format: str = None, default_profile: str = None, default_device: str = None) -> tuple[str, str, str]:
        """Zobrazí elegantní dialog pro výběr formátu papíru, CPU/GPU a výkonu hardware."""
        dialog = tk.Toplevel(self.root)
        dialog.title("Nastavení zpracování výkresů")
        dialog.geometry("520x420")
        dialog.resizable(False, False)
        dialog.grab_set()
        dialog.attributes("-topmost", True)
        dialog.focus_force()

        # 1. Formát papíru
        tk.Label(dialog, text="1. Formát papíru pro název souboru:", font=("Arial", 10, "bold")).pack(anchor="w", padx=20, pady=(12, 2))
        def_format = default_format or self.config.get("paper_format", "A4")
        format_var = tk.StringVar(value=def_format)
        format_entry = tk.Entry(dialog, textvariable=format_var, font=("Arial", 10), width=15)
        format_entry.pack(anchor="w", padx=20, pady=(0, 8))

        # 2. Výběr výpočetního zařízení (CPU vs GPU) — why: docs slibuje auto|cuda|cpu, gui mělo jen cuda/cpu
        tk.Label(dialog, text="2. Výpočetní zařízení (Hardware):", font=("Arial", 10, "bold")).pack(anchor="w", padx=20, pady=(4, 2))
        detected_devices = get_detected_devices()
        saved_dev = default_device or self.config.get("device", "auto")
        if saved_dev not in ["auto", "cuda", "cpu", "xpu", "mps"]:
            saved_dev = "auto"
        elif saved_dev == "cuda" and not any(d[0] == "cuda" and d[2] for d in detected_devices):
            saved_dev = "auto"

        dev_var = tk.StringVar(value=saved_dev)

        # Auto radio — doporučeno
        f_auto = tk.Frame(dialog)
        f_auto.pack(fill="x", padx=20, pady=1)
        tk.Radiobutton(f_auto, text="🤖 Automaticky (doporučeno)", variable=dev_var, value="auto", font=("Arial", 9, "bold")).pack(anchor="w")
        tk.Label(f_auto, text="Vybere GPU pokud je dostupná, jinak CPU", font=("Arial", 8), fg="#555").pack(anchor="w", padx=22)

        for d_key, d_label, d_avail in detected_devices:
            state = "normal" if d_avail else "disabled"
            f = tk.Frame(dialog)
            f.pack(fill="x", padx=20, pady=1)
            rb = tk.Radiobutton(f, text=d_label, variable=dev_var, value=d_key, state=state,
                                font=("Arial", 9, "bold" if d_avail else "normal"))
            rb.pack(anchor="w")

        # 3. Úroveň využití hardware
        tk.Label(dialog, text="3. Výkon:", font=("Arial", 10, "bold")).pack(anchor="w", padx=20, pady=(10, 2))
        def_prof = default_profile or self.config.get("hardware_profile", "turbo")
        prof_var = tk.StringVar(value=def_prof)
        dyn_profiles = get_dynamic_hardware_profiles()

        for key, pinfo in dyn_profiles.items():
            f = tk.Frame(dialog)
            f.pack(fill="x", padx=20, pady=2)
            rb = tk.Radiobutton(f, text=pinfo["name"], variable=prof_var, value=key, font=("Arial", 9, "bold"))
            rb.pack(anchor="w")
            desc = tk.Label(f, text=pinfo["description"], font=("Arial", 8), fg="#555")
            desc.pack(anchor="w", padx=22)

        res = [None, None, None]

        def on_confirm():
            res[0] = format_var.get().strip() or "A4"
            res[1] = prof_var.get()
            res[2] = dev_var.get()
            dialog.destroy()

        btn_frame = tk.Frame(dialog)
        btn_frame.pack(fill="x", padx=20, pady=(15, 10))
        tk.Button(btn_frame, text="Spustit zpracování", bg="#1976D2", fg="white", font=("Arial", 10, "bold"),
                  command=on_confirm, width=18, pady=3).pack(side="right")
        tk.Button(btn_frame, text="Zrušit", font=("Arial", 10), command=lambda: dialog.destroy(), width=10).pack(side="right", padx=10)

        dialog.protocol("WM_DELETE_WINDOW", lambda: dialog.destroy())
        dialog.wait_window()
        return res[0], res[1], res[2]

    def process_files_thread(self):
        """Pracovní vlákno: načtení modelů a asynchronní pipeline zpracování výkresů."""
        # Příprava slovníku rozpracovaných/již načtených výsledků
        processed_dict = dict(self.cached_results) if hasattr(self, 'cached_results') and self.cached_results else {}
        total = len(self.files)

        files_to_process = [f for f in self.files if f not in processed_dict]

        if not files_to_process:
            # Všechny soubory již byly dříve zpracovány
            self.results = [(f, processed_dict[f].get("num"), processed_dict[f].get("conf")) for f in self.files if f in processed_dict]
            self.queue.put(("done", None))
            return

        def on_model_progress(text, pct):
            self.queue.put(("model_progress", (text, pct)))

        try:
            backend, reader, device = load_models(
                hardware_profile=self.hardware_profile,
                device_choice=self.device_choice,
                progress_callback=on_model_progress,
                on_fatal_error=lambda t, m: messagebox.showerror(t, m),
            )
        except Exception as e:
            logger.error("Chyba načítání modelů: %s", e, exc_info=True)
            self.queue.put(("error", f"Nepodařilo se inicializovat modely:\n{e}"))
            return

        self.queue.put(("models_ready", device))

        prof_info = get_dynamic_hardware_profiles().get(self.hardware_profile, {})
        max_dim = prof_info.get("max_image_dim", 1280)
        pause_s = prof_info.get("inter_file_pause_s", 0.0)
        io_workers = prof_info.get("io_workers", 4)
        # Device desync — why: po CUDA OOM fallback je device cpu, ale GUI drželo cuda dim 1280 místo 960/1024
        if device == "cpu" and max_dim > 1024:
            is_vp = getattr(backend, "is_visionpsy", False)
            new_dim = 768 if is_vp else 1024
            logger.info("Device fallback na CPU — sníženo max_image_dim z %d na %d", max_dim, new_dim)
            max_dim = new_dim
        elif device == "xpu" and max_dim > 1024:
            max_dim = 1024
            logger.info("Device XPU — sníženo max_image_dim na %d", max_dim)

        with ThreadPoolExecutor(max_workers=io_workers) as executor:
            paths = [os.path.join(self.folder, f) for f in files_to_process]

            # Prefetch: načte obrázek A připraví cropi najednou (překryje I/O s předchozí inference)
            def _prefetch(fpath, dim):
                img, img_resized = load_and_resize_image(fpath, dim)
                crops = prepare_crops(img, device)
                return img, img_resized, crops

            next_future = executor.submit(_prefetch, paths[0], max_dim)

            for idx, fname in enumerate(files_to_process):
                # Kontrola požadavku na pozastavení před začátkem souboru
                if self._pause_requested.is_set():
                    logger.info("Zpracování pozastaveno uživatelem před souborem %s", fname)
                    self.queue.put(("paused", len(processed_dict)))
                    return

                done_count = len(processed_dict)
                pct = int((done_count / total) * 100)
                self.queue.put(("file_progress", (done_count, total, fname, pct, device)))

                fpath = paths[idx]
                try:
                    img, img_resized, prepped_crops = next_future.result()

                    if idx + 1 < len(files_to_process):
                        next_future = executor.submit(_prefetch, paths[idx + 1], max_dim)

                    final_num, conf_desc, vlm_ans, ocr_text = process_image(
                        img, img_resized, backend, reader, device,
                        prepped_crops=prepped_crops
                    )
                    logger.info("Výsledek %s -> číslo: %s, spolehlivost: %s", fname, final_num, conf_desc)
                    is_ver = str(conf_desc).startswith("Ověřeno")
                    processed_dict[fname] = {
                        "num": final_num,
                        "conf": str(conf_desc),
                        "user_num": clean_only_drawing_number(final_num) if final_num else "",
                        "user_name": "",
                        "is_verified": is_ver,
                        "is_checked": is_ver
                    }

                    # Průběžné uložení checkpointu po každém souboru
                    save_folder_checkpoint(self.folder, {
                        "folder": self.folder,
                        "timestamp": time.time(),
                        "updated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
                        "paper_format": self.paper_format,
                        "hardware_profile": self.hardware_profile,
                        "device": self.device_choice,
                        "total_files": total,
                        "status": "processing",
                        "results": processed_dict
                    })

                    if pause_s > 0:
                        time.sleep(pause_s)

                except Exception as e:
                    logger.error("Chyba zpracování %s: %s", fname, e, exc_info=True)
                    processed_dict[fname] = {
                        "num": None,
                        "conf": f"Chyba: {e}",
                        "user_num": "",
                        "user_name": "",
                        "is_verified": False,
                        "is_checked": False
                    }
                    save_folder_checkpoint(self.folder, {
                        "folder": self.folder,
                        "timestamp": time.time(),
                        "updated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
                        "paper_format": self.paper_format,
                        "hardware_profile": self.hardware_profile,
                        "device": self.device_choice,
                        "total_files": total,
                        "status": "processing",
                        "results": processed_dict
                    })
                    if idx + 1 < len(files_to_process):
                        next_future = executor.submit(_prefetch, paths[idx + 1], max_dim)

                # Kontrola pozastavení po dokončení souboru
                if self._pause_requested.is_set():
                    logger.info("Zpracování pozastaveno uživatelem po souboru %s", fname)
                    self.queue.put(("paused", len(processed_dict)))
                    return

        # Všechny soubory dokončeny
        save_folder_checkpoint(self.folder, {
            "folder": self.folder,
            "timestamp": time.time(),
            "updated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "paper_format": self.paper_format,
            "hardware_profile": self.hardware_profile,
            "device": self.device_choice,
            "total_files": total,
            "status": "review",
            "results": processed_dict
        })
        self.results = [(f, processed_dict[f].get("num"), processed_dict[f].get("conf")) for f in self.files if f in processed_dict]
        self.queue.put(("done", None))

    def check_queue(self):
        """Aktualizuje UI na základě zpráv z pracovního vlákna."""
        should_continue = True
        try:
            while True:
                msg_type, data = self.queue.get_nowait()
                if msg_type == "error":
                    messagebox.showerror("Chyba", data)
                    self.on_exit()
                    return
                elif msg_type == "model_progress":
                    text, pct = data
                    self.lbl_status.config(text=text)
                    self.lbl_pct.config(text=f"{pct}%")
                elif msg_type == "models_ready":
                    self.progress.stop()
                    self.progress.config(mode="determinate", maximum=100, value=0)
                    self._processing_start_time = time.time()
                    self._file_times = []
                    self._last_file_time = time.time()
                elif msg_type == "file_progress":
                    i, total, fname, pct, dev = data
                    self.lbl_status.config(text=f"Zpracovávám: {fname}")
                    # ETA výpočet — why: lineární průměr přestřeloval po warmupu (49m vs 8m), použít klouzavý průměr posledních 5
                    eta_text = ""
                    if self._processing_start_time and i > 0:
                        now = time.time()
                        if hasattr(self, '_last_file_time'):
                            try:
                                dt = now - self._last_file_time
                                if 0 < dt < 300:
                                    self._file_times.append(dt)
                                    if len(self._file_times) > 5:
                                        self._file_times.pop(0)
                                self._last_file_time = now
                            except Exception:
                                pass
                        if getattr(self, '_file_times', None) and len(self._file_times) >= 2:
                            avg_per_file = sum(self._file_times) / len(self._file_times)
                        else:
                            elapsed = now - self._processing_start_time
                            avg_per_file = elapsed / max(1, i)
                        remaining = (total - i) * avg_per_file
                        if remaining >= 60:
                            eta_text = f" | Zbývá ~{int(remaining // 60)} min {int(remaining % 60)} s"
                        else:
                            eta_text = f" | Zbývá ~{int(remaining)} s"
                        eta_text += f" ({avg_per_file:.1f} s/soubor)"
                    self.lbl_pct.config(text=f"{pct}%{eta_text}")
                    prof_lbl = get_dynamic_hardware_profiles().get(self.hardware_profile, {}).get("short_name", "")
                    self.lbl_counter.config(text=f"{i + 1} / {total} souborů | {prof_lbl} | {dev.upper()}")
                    self.progress["value"] = pct
                elif msg_type == "paused":
                    should_continue = False
                    count_done = data
                    messagebox.showinfo(
                        "Zpracování pozastaveno",
                        f"Postup byl bezpečně uložen ({count_done} z {len(self.files)} výkresů hotovo).\n\n"
                        f"Počítač můžete bez obav vypnout.\n"
                        f"Při příštím spuštění programu a zvolení této složky budete moci plynule navázat."
                    )
                    self.on_exit()
                    return
                elif msg_type == "done":
                    should_continue = False
                    self.show_review_gui()
                    return
        except queue.Empty:
            pass
        except Exception as e:
            logger.error("Chyba v check_queue UI: %s", e, exc_info=True)
        
        if should_continue:
            self.root.after(100, self.check_queue)


    def _update_stats_header(self):
        """Aktualizuje počitadlo ověřených a zbývajících výkresů v horní liště."""
        if not hasattr(self, 'row_items') or not self.row_items:
            return
        verified_cnt = sum(1 for item in self.row_items if item.get('is_verified') or item['var'].get())
        review_cnt = len(self.row_items) - verified_cnt
        if hasattr(self, 'lbl_summary') and self.lbl_summary.winfo_exists():
            self.lbl_summary.config(
                text=f"Celkem: {len(self.row_items)} výkresů | ● Ověřeno: {verified_cnt} | ● K ověření: {review_cnt}   [Tažením příčky uprostřed lze měnit šířku panelů]"
            )

    def show_review_gui(self):
        """Zobrazí tabulku s navrženými názvy, spolehlivostí a bočním živým náhledem v nastavitelném Splitteru (PanedWindow)."""
        for w in list(self.root.winfo_children()):
            try:
                w.destroy()
            except Exception:
                pass

        self.root.geometry("1600x920")
        try:
            self.root.state('zoomed')
        except Exception:
            pass
        self.root.title("Kontrola a přejmenování výkresů")
        self.root.protocol("WM_DELETE_WINDOW", self.on_review_window_close)

        self._zoom = 1.0  # Aktuální zoom celého výkresu (1.0 = fit)
        self._pan_x = 0
        self._pan_y = 0
        self._pil_orig = None  # Originál PIL pro zoom bez ztráty kvality
        self.preview_cache = OrderedDict()
        self.current_preview_fname = None

        # Bezpečné rozbalení výsledků (podpora pro 2-, 3- i 4-prvkové tuples)
        def _normalize_item(item):
            if len(item) >= 4:
                return item[0], item[1], item[3]
            elif len(item) == 3:
                return item[0], item[1], item[2]
            elif len(item) == 2:
                return item[0], item[1], "Nenalezeno"
            return item[0], None, "Nenalezeno"

        # Počáteční řazení: Všechny ČERVENÉ (k ověření) nahoru, ZELENÉ (ověřené) dolů
        def _get_sort_priority(item):
            norm = _normalize_item(item)
            fname, num, conf = norm
            if str(conf).startswith("Ověřeno"):
                return (1, fname)
            return (0, fname)

        sorted_results = [_normalize_item(it) for it in sorted(self.results, key=_get_sort_priority)]

        # Statistika spolehlivosti pro horní lištu
        verified_cnt = sum(1 for _, _, c in sorted_results if str(c).startswith("Ověřeno"))
        review_cnt = len(sorted_results) - verified_cnt

        summary_frame = tk.Frame(self.root, bg="#ECEFF1", pady=6)
        summary_frame.pack(fill=tk.X)

        # Režim ověřování: True = auto (AI+OCR shoda = zelená), False = manuální (vše červené)
        if not hasattr(self, '_auto_verify_mode'):
            self._auto_verify_mode = True

        self.lbl_summary = tk.Label(
            summary_frame,
            text=f"Celkem: {len(sorted_results)} výkresů | ● Ověřeno (dole): {verified_cnt} | ● K ověření (nahoře): {review_cnt}   [Tažením příčky uprostřed lze měnit šířku panelů]",
            font=("Segoe UI", 11, "bold"),
            bg="#ECEFF1",
            fg="#263238"
        )
        self.lbl_summary.pack(side=tk.LEFT, fill=tk.X, expand=True)

        def _toggle_verify_mode():
            self._auto_verify_mode = not self._auto_verify_mode
            mode_label = "Auto-ověření" if self._auto_verify_mode else "Manuální kontrola"
            btn_mode.config(text=f"⚙ {mode_label}")
            for r in self.row_items:
                # Výkresy ověřené uživatelem (Enterem) zůstanou zelené vždy
                if r.get('conf') == "Ověřeno (Uživatelem)":
                    continue
                if self._auto_verify_mode:
                    # Obnovit původní stav podle AI+OCR konsensu
                    is_ver = str(r.get('original_conf', r['conf'])).startswith("Ověřeno")
                    r['conf'] = r.get('original_conf', r['conf'])
                else:
                    # Vše na červenou – manuální kontrola
                    is_ver = False
                r['is_verified'] = is_ver
                r['var'].set(is_ver)
                new_color = "#4CAF50" if is_ver else "#F44336"
                r['lbl_status'].delete("all")
                r['lbl_status'].create_oval(3, 3, 15, 15, fill=new_color, outline=new_color)
            # Přeřadit: červené nahoru, zelené dolů
            self.row_items.sort(key=lambda x: 1 if x['is_verified'] else 0)
            _relayout_rows()
            self._update_stats_header()

        mode_label = "Auto-ověření" if self._auto_verify_mode else "Manuální kontrola"
        btn_mode = tk.Button(
            summary_frame, text=f"⚙ {mode_label}",
            font=("Segoe UI", 9, "bold"), bg="#546E7A", fg="white",
            activebackground="#455A64", activeforeground="white",
            relief="flat", padx=10, pady=3,
            command=_toggle_verify_mode
        )
        btn_mode.pack(side=tk.RIGHT, padx=10)

        # --- HLAVNÍ NASTAVITELNÝ SPLITTER (PanedWindow) ---
        paned = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        # --- LEVÝ PANEL: Tabulka výkresů s vertikálním i horizontálním posuvem ---
        left_container = tk.Frame(paned, bg="#F5F5F5")
        paned.add(left_container, weight=1)

        canvas = tk.Canvas(left_container, highlightthickness=0, bg="#FFFFFF")
        v_scrollbar = ttk.Scrollbar(left_container, orient="vertical", command=canvas.yview)
        h_scrollbar = ttk.Scrollbar(left_container, orient="horizontal", command=canvas.xview)
        scroll_frame = tk.Frame(canvas, bg="#FFFFFF")

        scroll_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        canvas_window = canvas.create_window((0, 0), window=scroll_frame, anchor="nw")
        
        def _on_canvas_configure(event):
            canvas.itemconfig(canvas_window, width=max(event.width, scroll_frame.winfo_reqwidth()))
        canvas.bind("<Configure>", _on_canvas_configure)

        canvas.configure(xscrollcommand=h_scrollbar.set, yscrollcommand=v_scrollbar.set)

        # Layout scrollbarů a canvasu
        canvas.grid(row=0, column=0, sticky="nsew")
        v_scrollbar.grid(row=0, column=1, sticky="ns")
        h_scrollbar.grid(row=1, column=0, sticky="ew")
        left_container.rowconfigure(0, weight=1)
        left_container.columnconfigure(0, weight=1)

        # Rolování myší – aktivuje se pouze nad levým panelem (Enter/Leave pattern)
        def _bind_mouse_scroll(e):
            canvas.bind_all("<MouseWheel>", lambda ev: canvas.yview_scroll(-1 * (ev.delta // 120), "units"))
            canvas.bind_all("<Shift-MouseWheel>", lambda ev: canvas.xview_scroll(-1 * (ev.delta // 120), "units"))
        def _unbind_mouse_scroll(e):
            canvas.unbind_all("<MouseWheel>")
            canvas.unbind_all("<Shift-MouseWheel>")
        left_container.bind("<Enter>", _bind_mouse_scroll)
        left_container.bind("<Leave>", _unbind_mouse_scroll)

        # --- PRAVÝ PANEL: Dominantní živý náhled (vždy celý výkres + zoom kolečkem) ---
        preview_frame = tk.LabelFrame(paned, text="🔍 Živý náhled výkresu", font=("Segoe UI", 11, "bold"), padx=8, pady=8)
        paned.add(preview_frame, weight=2)

        top_prev_bar = tk.Frame(preview_frame)
        top_prev_bar.pack(fill=tk.X, pady=(0, 6))

        self.lbl_preview_name = tk.Label(top_prev_bar, text="Vyberte výkres", font=("Segoe UI", 11, "bold"), fg="#1565C0", anchor="w")
        self.lbl_preview_name.pack(side=tk.LEFT, fill=tk.X, expand=True)

        self.btn_preview_mode = tk.Button(top_prev_bar, text="⟲ 100% (fit)", font=("Segoe UI", 9), command=self._reset_zoom)
        self.btn_preview_mode.pack(side=tk.RIGHT, padx=4)

        tk.Button(top_prev_bar, text="Otevřít (F3)", font=("Segoe UI", 9), command=lambda: self.open_file(self.current_preview_fname)).pack(side=tk.RIGHT)

        # Canvas místo Label – why: umí scroll/pan při zoom > fit a kolečko zoom
        self.preview_canvas = tk.Canvas(preview_frame, bg="#1E272C", highlightthickness=0)
        self.preview_canvas.pack(fill=tk.BOTH, expand=True, pady=4)
        self.preview_canvas.bind("<Configure>", self._on_preview_resize)
        # zoom kolečkem (Win/Mac delta, Linux Button-4/5)
        self.preview_canvas.bind("<MouseWheel>", self._on_zoom)
        self.preview_canvas.bind("<Button-4>", lambda e: self._on_zoom(e, delta=120))
        self.preview_canvas.bind("<Button-5>", lambda e: self._on_zoom(e, delta=-120))
        # pan tažením myši při přiblížení
        self.preview_canvas.bind("<ButtonPress-1>", self._on_pan_start)
        self.preview_canvas.bind("<B1-Motion>", self._on_pan_move)
        self._canvas_img_id = None

        # TUČNÉ ROZPOZNANÉ ČÍSLO VYCENTROVANÉ PŘÍMO POD OBRÁZKEM
        bottom_info_card = tk.Frame(preview_frame, bg="#263238", pady=8, padx=16)
        bottom_info_card.pack(fill=tk.X, pady=(4, 0))

        self.lbl_preview_detected_num = tk.Label(
            bottom_info_card,
            text="📌 Rozpoznané číslo: ---",
            font=("Segoe UI", 17, "bold"),
            bg="#263238",
            fg="#FFD54F",  # Zářivě žlutá pro okamžitou viditelnost
            anchor="center",
            justify="center"
        )
        self.lbl_preview_detected_num.pack(fill=tk.X, expand=True, pady=(1, 1))

        self.lbl_preview_conf = tk.Label(
            bottom_info_card,
            text="",
            font=("Segoe UI", 10, "bold"),
            bg="#263238",
            fg="#81C784",
            anchor="center",
            justify="center"
        )
        self.lbl_preview_conf.pack(fill=tk.X, expand=True, pady=(1, 1))

        # Sloupce tabulky: 0: Soubor, 1: Číslo výkresu, 2: Název výkresu, 3: Stav (tečka), 4: Přejmenovat
        scroll_frame.columnconfigure(0, weight=2)
        scroll_frame.columnconfigure(1, weight=2)
        scroll_frame.columnconfigure(2, weight=3)
        scroll_frame.columnconfigure(3, weight=0)
        scroll_frame.columnconfigure(4, weight=0)

        headers = [
            ("Původní soubor", 20), ("Číslo výkresu", 16),
            ("Název výkresu", 24), ("Stav", 6),
            ("Přejmenovat?", 12),
        ]
        for col, (text, width) in enumerate(headers):
            tk.Label(
                scroll_frame, text=text, font=("Segoe UI", 9, "bold"),
                bg="#ECEFF1", fg="#37474F", width=width,
                anchor="w" if col < 3 else "center", padx=4, pady=4
            ).grid(row=0, column=col, sticky="ew", padx=2, pady=4)

        # Inicializace datových objektů pro každý řádek (jednorázový layout)
        self.row_items = []

        # Obnovení uložených stavů z checkpointu, pokud existují
        saved_items_map = {}
        checkpoint_data = get_folder_checkpoint(self.folder)
        if checkpoint_data and "results" in checkpoint_data:
            saved_items_map = checkpoint_data["results"]

        for idx, (fname, num, conf) in enumerate(sorted_results):
            row_num = idx + 1
            lbl_fn = tk.Label(scroll_frame, text=fname, anchor="w", cursor="hand2", font=("Segoe UI", 9), bg="#FFFFFF")
            lbl_fn.grid(row=row_num, column=0, sticky="ew", padx=3, pady=2)

            ext = os.path.splitext(fname)[1]

            # Kontrola uložených uživatelských úprav z předchozí relace
            saved_item = saved_items_map.get(fname, {})
            if saved_item.get("user_num") is not None and str(saved_item.get("user_num")).strip() != "":
                init_num = str(saved_item.get("user_num")).strip()
            else:
                init_num = clean_only_drawing_number(num) if num else ""

            init_name = str(saved_item.get("user_name") or saved_item.get("user_title") or "").strip()

            if "is_verified" in saved_item:
                is_verified = bool(saved_item["is_verified"])
            else:
                is_verified = str(conf).startswith("Ověřeno")

            if "is_checked" in saved_item:
                is_checked = bool(saved_item["is_checked"])
            else:
                is_checked = is_verified

            if saved_item.get("conf"):
                conf = saved_item.get("conf")

            entry_num = tk.Entry(scroll_frame, font=("Segoe UI", 9, "bold"))
            entry_num.insert(0, init_num)
            entry_num.grid(row=row_num, column=1, sticky="ew", padx=3, pady=2)

            entry_name = tk.Entry(scroll_frame, font=("Segoe UI", 9))
            entry_name.insert(0, init_name)
            entry_name.grid(row=row_num, column=2, sticky="ew", padx=3, pady=2)

            dot_color = "#4CAF50" if is_verified else "#F44336"
            status_canvas = tk.Canvas(scroll_frame, width=18, height=18, bg="#FFFFFF", highlightthickness=0)
            status_canvas.create_oval(3, 3, 15, 15, fill=dot_color, outline=dot_color)
            status_canvas.grid(row=row_num, column=3, sticky="", padx=3, pady=2)

            var = tk.BooleanVar(value=is_checked)
            chk_widget = tk.Checkbutton(scroll_frame, variable=var, bg="#FFFFFF")
            chk_widget.grid(row=row_num, column=4, sticky="ew", padx=3, pady=2)

            # Formát vždy jen z volby uživatele na startu — why: Excel/rename nesmí brát formát z názvu souboru
            r_data = {
                'idx': idx,
                'fname': fname,
                'num': num,
                'conf': conf,
                'original_conf': conf,
                'is_verified': is_verified,
                'lbl_fn': lbl_fn,
                'entry_num': entry_num,
                'entry_name': entry_name,
                'lbl_status': status_canvas,
                'var': var,
                'chk': chk_widget,
                'ext': ext,
                'format': self.paper_format
            }
            self.row_items.append(r_data)

        total_cnt = len(self.row_items)

        def _ensure_row_visible(idx_val):
            if total_cnt <= 1:
                return
            fraction = max(0.0, min(1.0, max(0, idx_val - 2) / max(1, total_cnt)))
            canvas.yview_moveto(fraction)

        def _relayout_rows():
            """Znovu rozloží všechny řádky do mřížky podle aktuálního pořadí row_items."""
            for j, rr in enumerate(self.row_items):
                grid_row = j + 1
                rr['lbl_fn'].grid(row=grid_row, column=0, sticky="ew", padx=3, pady=2)
                rr['entry_num'].grid(row=grid_row, column=1, sticky="ew", padx=3, pady=2)
                rr['entry_name'].grid(row=grid_row, column=2, sticky="ew", padx=3, pady=2)
                rr['lbl_status'].grid(row=grid_row, column=3, sticky="ew", padx=3, pady=2)
                rr['chk'].grid(row=grid_row, column=4, sticky="ew", padx=3, pady=2)

        def _confirm_and_move_to_verified(idx, focus_target="entry_num"):
            """Označí výkres Enterem jako ověřený a fyzicky ho přesune dolů do zelené sekce."""
            if idx < 0 or idx >= len(self.row_items):
                return
            r = self.row_items[idx]
            r['var'].set(True)
            r['is_verified'] = True
            r['conf'] = "Ověřeno (Uživatelem)"
            r['lbl_status'].delete("all")
            r['lbl_status'].create_oval(3, 3, 15, 15, fill="#4CAF50", outline="#4CAF50")

            # Přesun dolů: ověřené výkresy patří na konec seznamu (zelená sekce)
            self.row_items.sort(key=lambda x: 1 if x['is_verified'] else 0)
            _relayout_rows()
            self._update_stats_header()

            # Fokus na položku, která je nyní na stejném místě (původně následující)
            next_i = min(idx, len(self.row_items) - 1)
            if next_i >= 0:
                next_row = self.row_items[next_i]
                target_widget = next_row['entry_name'] if focus_target == "entry_name" else next_row['entry_num']
                target_widget.focus_set()
                target_widget.select_range(0, tk.END)
                _ensure_row_visible(next_i)
                self.update_live_preview(
                    next_row['fname'],
                    next_row['conf'],
                    next_row['entry_num'].get().strip()
                )

        # Jednorázové navázání událostí na všechny řádky
        for i, r in enumerate(self.row_items):
            def _bind_row(curr_i=i, row_data=r):
                fn = row_data['fname']
                ent_n = row_data['entry_num']
                ent_title = row_data['entry_name']
                lfn = row_data['lbl_fn']

                def _select_this_row():
                    self.update_live_preview(fn, row_data['conf'], ent_n.get().strip())

                def _navigate(direction, field):
                    """Sdílená navigace šipkami nahoru/dolů pro oba sloupce."""
                    cur = self.row_items.index(row_data)
                    target_i = max(0, cur - 1) if direction == "up" else min(len(self.row_items) - 1, cur + 1)
                    target_r = self.row_items[target_i]
                    target_r[field].focus_set()
                    target_r[field].select_range(0, tk.END)
                    _ensure_row_visible(target_i)
                    self.update_live_preview(target_r['fname'], target_r['conf'], target_r['entry_num'].get().strip())
                    return "break"

                def _confirm(field):
                    """Sdílené potvrzení Enterem pro oba sloupce."""
                    _confirm_and_move_to_verified(self.row_items.index(row_data), focus_target=field)
                    return "break"

                def _on_key_release(e=None):
                    clean_n = clean_only_drawing_number(ent_n.get().strip())
                    self.lbl_preview_detected_num.config(text=f"📌 Rozpoznané číslo: {clean_n}")

                def _handle_preview(e=None):
                    self.open_file(fn)
                    return "break"

                # Bindings pro oba sloupce (entry_num a entry_name)
                lfn.bind("<Button-1>", lambda e: (ent_n.focus_set(), _select_this_row(), "break")[-1])
                for field_key, widget in [("entry_num", ent_n), ("entry_name", ent_title)]:
                    widget.bind("<Up>", lambda e, f=field_key: _navigate("up", f))
                    widget.bind("<Down>", lambda e, f=field_key: _navigate("down", f))
                    widget.bind("<Return>", lambda e, f=field_key: _confirm(f))
                    widget.bind("<KP_Enter>", lambda e, f=field_key: _confirm(f))
                    widget.bind("<Control-o>", _handle_preview)
                    widget.bind("<F3>", _handle_preview)
                    widget.bind("<FocusIn>", lambda e: _select_this_row())
                ent_n.bind("<KeyRelease>", _on_key_release)

            _bind_row()

        # Nastavení počátečního fokusu
        if self.row_items:
            first_row = self.row_items[0]
            first_row['entry_num'].focus_set()
            first_row['entry_num'].select_range(0, tk.END)
            self.update_live_preview(
                first_row['fname'],
                first_row['conf'],
                first_row['entry_num'].get().strip()
            )

        # Nastavení výchozí šířky panelu Splitteru po vykreslení
        try:
            self.root.update_idletasks()
            paned.sashpos(0, 720)
        except Exception:
            pass

        btn_frame = tk.Frame(self.root, pady=10)
        btn_frame.pack(fill=tk.X)

        tk.Button(btn_frame, text="✓ Vybrat vše (Ctrl+A)", command=lambda: self._set_all_checks(True)).pack(side=tk.LEFT, padx=10)
        tk.Button(btn_frame, text="✗ Odškrtnout vše (Ctrl+D)", command=lambda: self._set_all_checks(False)).pack(side=tk.LEFT, padx=5)

        tk.Button(btn_frame, text="Přejmenovat vybrané (Ctrl+Enter)", bg="#4CAF50", fg="white", font=("Arial", 10, "bold"), command=self.rename_files).pack(side=tk.RIGHT, padx=10)

        # Globální klávesové zkratky
        self.root.bind("<Control-a>", lambda e: self._set_all_checks(True))
        self.root.bind("<Control-d>", lambda e: self._set_all_checks(False))
        self.root.bind("<Control-Return>", lambda e: self.rename_files())

    def _save_current_review_state(self):
        """Uloží kompletní stav revizního okna (uživatelské úpravy čísel, názvů, ověření, checkboxy)."""
        if not hasattr(self, 'row_items') or not self.row_items or not self.folder:
            return
        review_dict = {}
        for item in self.row_items:
            fn = item['fname']
            num_val = item['entry_num'].get().strip()
            name_val = item['entry_name'].get().strip()
            conf_val = item['conf']
            is_v = item.get('is_verified', False)
            is_c = item['var'].get()
            review_dict[fn] = {
                "num": item.get('num'),
                "conf": str(conf_val),
                "user_num": num_val,
                "user_name": name_val,
                "is_verified": is_v,
                "is_checked": is_c
            }
        save_folder_checkpoint(self.folder, {
            "folder": self.folder,
            "timestamp": time.time(),
            "updated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "paper_format": getattr(self, 'paper_format', 'A4'),
            "hardware_profile": getattr(self, 'hardware_profile', 'turbo'),
            "device": getattr(self, 'device_choice', 'auto'),
            "total_files": len(self.files) if hasattr(self, 'files') else len(review_dict),
            "status": "review",
            "results": review_dict
        })
        logger.info("Uložen kompletní stav revizního okna do checkpointu (%d položek)", len(review_dict))

    def export_all_items_to_excel(self) -> bool:
        """
        Vyexportuje všechny výkresy z tabulky do Excelu 'seznam_vykresu.xlsx' ve vybrané složce.
        Sloupce: Číslo výkresu, Název výkresu, Formát.
        """
        if not hasattr(self, 'row_items') or not self.row_items or not self.folder:
            return False

        export_items = []
        for item in self.row_items:
            fname = item['fname']
            num_val = clean_only_drawing_number(item['entry_num'].get().strip())
            if not num_val:
                num_val = clean_only_drawing_number(fname)
            name_val = item['entry_name'].get().strip()
            # Formát vždy jen z volby uživatele na startu (self.paper_format)
            export_items.append((num_val, name_val, self.paper_format))

        if not export_items:
            return False

        excel_path = os.path.join(self.folder, "seznam_vykresu.xlsx")
        return export_to_excel(excel_path, export_items, self.paper_format)

    def on_review_window_close(self):
        """Při zavření revizního okna křížkem automaticky uloží stav, vyexportuje do Excelu a ukončí aplikaci."""
        try:
            self._save_current_review_state()
            ok = self.export_all_items_to_excel()
            if not ok:
                logger.warning("Export do Excelu při zavření vrátil False")
        except Exception as e:
            logger.error("Chyba při automatickém ukládání na konci: %s", e)
            try:
                messagebox.showwarning("Varování", f"Nepodařilo se uložit Excel seznam_vykresu.xlsx: {e}\nZkontrolujte práva k zápisu ve složce.")
            except Exception:
                pass
        self.on_exit()

    def _reset_zoom(self):
        """Vrátí zoom na fit (100%) a vycentruje."""
        self._zoom = 1.0
        self._pan_x = 0
        self._pan_y = 0
        if self.current_preview_fname:
            self.update_live_preview(self.current_preview_fname, force_reload=True)

    def _on_zoom(self, event, delta=None):
        """Kolečko myši – why: klasický zoom 0.25x až 5x, centered."""
        if self._pil_orig is None:
            return
        d = delta if delta is not None else getattr(event, 'delta', 0)
        # Linux Button-4/5 už posílá delta
        if d > 0:
            factor = 1.18
        elif d < 0:
            factor = 1 / 1.18
        else:
            return
        new_zoom = max(0.25, min(5.0, self._zoom * factor))
        if abs(new_zoom - self._zoom) < 0.01:
            return
        self._zoom = new_zoom
        self._render_zoomed()
        return "break"

    def _on_pan_start(self, event):
        self._pan_start_x = event.x - self._pan_x
        self._pan_start_y = event.y - self._pan_y

    def _on_pan_move(self, event):
        if self._zoom <= 1.0:
            return
        self._pan_x = event.x - self._pan_start_x
        self._pan_y = event.y - self._pan_start_y
        self._render_zoomed()
        return "break"

    def _on_preview_resize(self, event=None):
        """Při změně velikosti okna překreslí fit (jen pokud zoom==1.0)."""
        if self._zoom == 1.0 and self.current_preview_fname:
            # debounce – necháme idle, aby neblikalo při tažení sash
            self.root.after_idle(lambda: self.update_live_preview(self.current_preview_fname, force_reload=True))

    @staticmethod
    def _process_preview_image(pil_img: Image.Image, tw: int, th: int, zoom: float = 1.0) -> Image.Image:
        """Shared PIL pipeline: grayscale → autocontrast → fit+zoom resize → sharpen."""
        render_img = pil_img.convert("L")
        render_img = ImageOps.autocontrast(render_img, cutoff=1)
        w_orig, h_orig = render_img.size
        base_scale = min(tw / w_orig, th / h_orig)
        scale = base_scale * zoom
        nw = max(1, int(w_orig * scale))
        nh = max(1, int(h_orig * scale))
        render_img = render_img.resize((nw, nh), Image.Resampling.BILINEAR)
        return ImageEnhance.Sharpness(render_img).enhance(1.35)

    def _render_image_background(self, file_p: str, tw: int, th: int, zoom: float = 1.0):
        """Pomocná funkce pro asynchronní načtení – vždy celý výkres, zoom = násobek fit."""
        try:
            with Image.open(file_p) as img:
                render_img = self._process_preview_image(img, tw, th, zoom)
                return render_img.convert("RGB"), None
        except Exception as e:
            return None, e

    def _render_zoomed(self):
        """Překreslí aktuální PIL originál se současným zoom/pan do Canvas."""
        if self._pil_orig is None or not hasattr(self, 'preview_canvas'):
            return
        try:
            cw = self.preview_canvas.winfo_width()
            ch = self.preview_canvas.winfo_height()
            tw = max(600, cw - 16) if cw > 50 else 1000
            th = max(500, ch - 16) if ch > 50 else 800
            render_img = self._process_preview_image(self._pil_orig, tw, th, self._zoom)
            rgb = render_img.convert("RGB")
            photo = ImageTk.PhotoImage(rgb)
            self.preview_canvas.delete("all")
            nw, nh = rgb.size
            cx = cw // 2 + self._pan_x if cw > 50 else nw // 2
            cy = ch // 2 + self._pan_y if ch > 50 else nh // 2
            self._canvas_img_id = self.preview_canvas.create_image(cx, cy, image=photo, anchor="center")
            self._current_photo = photo
            self.preview_canvas.image = photo
            self.btn_preview_mode.config(text=f"⟲ {int(self._zoom*100)}%")
        except Exception as e:
            logger.error("Chyba _render_zoomed: %s", e)

    def update_live_preview(self, fname: str, conf_desc: str = "", detected_num: str = "", force_reload: bool = False):
        """Okamžitě aktualizuje textové info (0ms) a asynchronně načte celý výkres – vždy full, crop odstraněn."""
        if not fname or not self.folder:
            return
        # při přepnutí souboru reset zoom/pan, pokud to není force reload téhož souboru
        if fname != self.current_preview_fname:
            self._zoom = 1.0
            self._pan_x = 0
            self._pan_y = 0
        self.current_preview_fname = fname
        self.lbl_preview_name.config(text=f"📄 {fname}")

        # 1. Okamžitá aktualizace textů pod náhledem
        if not detected_num and hasattr(self, 'row_items'):
            for item in self.row_items:
                if item['fname'] == fname:
                    detected_num = item['entry_num'].get().strip()
                    break

        if hasattr(self, 'lbl_preview_detected_num'):
            clean_display = clean_only_drawing_number(detected_num or fname)
            self.lbl_preview_detected_num.config(text=f"📌 Rozpoznané číslo: {clean_display}")

        if hasattr(self, 'lbl_preview_conf'):
            if conf_desc:
                if conf_desc.startswith("Ověřeno"):
                    self.lbl_preview_conf.config(text=f"● {conf_desc}", fg="#4CAF50")
                else:
                    self.lbl_preview_conf.config(text=f"● {conf_desc}", fg="#F44336")
            else:
                self.lbl_preview_conf.config(text="")

        # 2. Kontrola cache (jen fit 100% se cachuje)
        cache_key = fname
        if not force_reload and self._zoom == 1.0 and cache_key in self.preview_cache:
            # cache obsahuje fit PhotoImage – obnovíme i _pil_orig pokud chybí
            photo = self.preview_cache[cache_key]
            # pro zoom potřebujeme originál – pokud chybí, načteme znova
            if self._pil_orig is None:
                try:
                    with Image.open(os.path.join(self.folder, fname)) as im:
                        self._pil_orig = im.copy()
                except Exception:
                    pass
            self.preview_canvas.delete("all")
            cw = self.preview_canvas.winfo_width()
            ch = self.preview_canvas.winfo_height()
            cx = cw // 2 if cw > 50 else 0
            cy = ch // 2 if ch > 50 else 0
            self.preview_canvas.create_image(cx, cy, image=photo, anchor="center")
            self.preview_canvas.image = photo
            self._current_photo = photo
            self.btn_preview_mode.config(text=f"⟲ {int(self._zoom*100)}%")
            return

        full_path = os.path.join(self.folder, fname)
        if not os.path.isfile(full_path):
            return

        cw = self.preview_canvas.winfo_width()
        ch = self.preview_canvas.winfo_height()
        target_w = max(600, cw - 16) if cw > 50 else 1000
        target_h = max(500, ch - 16) if ch > 50 else 800

        self._preview_request_id += 1
        req_id = self._preview_request_id
        cur_zoom = self._zoom

        def _worker():
            # načti originál pro další zoom bez ztráty kvality
            try:
                with Image.open(full_path) as im:
                    pil_copy = im.copy()
            except Exception as e:
                logger.error("Chyba načítání %s: %s", fname, e)
                return
            rgb_img, err = self._render_image_background(full_path, target_w, target_h, zoom=cur_zoom)
            if rgb_img:
                def _apply():
                    try:
                        if req_id != self._preview_request_id:
                            return
                        self._pil_orig = pil_copy
                        photo = ImageTk.PhotoImage(rgb_img)
                        if cur_zoom == 1.0:
                            self.preview_cache[cache_key] = photo
                            while len(self.preview_cache) > _MAX_PREVIEW_CACHE_SIZE:
                                self.preview_cache.popitem(last=False)
                        self.preview_canvas.delete("all")
                        cw2 = self.preview_canvas.winfo_width()
                        ch2 = self.preview_canvas.winfo_height()
                        cx = cw2 // 2 + self._pan_x if cw2 > 50 else 0
                        cy = ch2 // 2 + self._pan_y if ch2 > 50 else 0
                        self.preview_canvas.create_image(cx, cy, image=photo, anchor="center")
                        self.preview_canvas.image = photo
                        self._current_photo = photo
                        self.btn_preview_mode.config(text=f"⟲ {int(self._zoom*100)}%")
                    except Exception as e:
                        logger.error("Chyba při aplikaci PhotoImage: %s", e)
                self.root.after_idle(_apply)
            elif err and req_id == self._preview_request_id:
                logger.error("Chyba načítání náhledu %s: %s", fname, err)

            # Prefetch dalšího souboru (fit only) pro okamžité přepnutí
            try:
                if hasattr(self, 'row_items'):
                    curr_idx = next((i for i, r in enumerate(self.row_items) if r['fname'] == fname), -1)
                    if 0 <= curr_idx < len(self.row_items) - 1:
                        next_fname = self.row_items[curr_idx + 1]['fname']
                        if next_fname not in self.preview_cache:
                            n_path = os.path.join(self.folder, next_fname)
                            n_img, _ = self._render_image_background(n_path, target_w, target_h, zoom=1.0)
                            if n_img:
                                def _cache_next(im=n_img, k=next_fname):
                                    try:
                                        self.preview_cache[k] = ImageTk.PhotoImage(im)
                                        while len(self.preview_cache) > _MAX_PREVIEW_CACHE_SIZE:
                                            self.preview_cache.popitem(last=False)
                                    except Exception:
                                        pass
                                self.root.after_idle(_cache_next)
            except Exception:
                pass

        self.preview_executor.submit(_worker)

    def _set_all_checks(self, state: bool):
        if hasattr(self, 'row_items'):
            for item in self.row_items:
                item['var'].set(state)
        self._update_stats_header()

    def open_file(self, fname: str):
        if not fname:
            return
        try:
            os.startfile(os.path.join(self.folder, fname))
        except Exception as e:
            logger.error("Chyba při otevírání náhledu: %s", e)
            messagebox.showerror("Chyba", f"Nepodařilo se otevřít soubor: {e}")

    def rename_files(self):
        # Spočítat kolik souborů bude přejmenováno
        to_rename = []
        for item in self.row_items:
            fname = item['fname']
            ent_num = item['entry_num']
            ent_name = item['entry_name']
            var = item['var']
            ext = item['ext']
            conf = item['conf']
            # Formát pro Excel vždy jen z volby uživatele na startu (self.paper_format)
            fmt_val = self.paper_format
            if not var.get():
                continue
            raw_val = ent_num.get().strip()
            if not raw_val:
                continue
            clean_num = clean_only_drawing_number(raw_val)
            if not clean_num:
                continue
            name_val = ent_name.get().strip()
            to_rename.append((fname, clean_num, name_val, ext, conf, fmt_val))

        if not to_rename:
            messagebox.showinfo("Informace", "Nejsou vybrány žádné soubory k přejmenování.")
            return

        # Potvrzovací dialog
        if not messagebox.askyesno(
            "Potvrzení přejmenování",
            f"Chystáte se přejmenovat {len(to_rename)} souborů.\n\n"
            f"Tato operace přejmenuje vybrané výkresy.\n"
            f"Záloha pro vrácení změn a Excel budou uloženy automaticky.\n\n"
            f"Pokračovat?",
            icon="warning"
        ):
            return

        undo_map = {}
        renamed_count = 0
        error_count = 0
        error_files = []
        verified_export_items = []

        successful_verified: list[tuple] = []
        successful_all: list[tuple] = []
        successful_fnames: set[str] = set()
        for fname, clean_num, name_val, ext, conf, fmt_val in to_rename:
            # Sanitizace zakázaných Windows znaků
            safe_num = sanitize_filename(clean_num)
            if not safe_num:
                error_count += 1
                error_files.append(f"{fname} (neplatné číslo po sanitizaci)")
                continue

            if self.paper_format:
                new_name = f"{safe_num} ({self.paper_format}){ext}"
            else:
                new_name = f"{safe_num}{ext}"

            src = os.path.join(self.folder, fname)
            dst = os.path.join(self.folder, new_name)

            success = False
            if src != dst:
                try:
                    if os.path.exists(dst):
                        base, ext_part = os.path.splitext(new_name)
                        # Windows styl: "135185 (A4) (2).tif", "135185 (A4) (3).tif"
                        counter = 2
                        while os.path.exists(os.path.join(self.folder, f"{base} ({counter}){ext_part}")):
                            counter += 1
                        dst = os.path.join(self.folder, f"{base} ({counter}){ext_part}")

                    os.rename(src, dst)
                    undo_map[os.path.basename(dst)] = fname
                    renamed_count += 1
                    success = True

                    # Průběžný zápis undo zálohy — atomický per-folder
                    for _undo_path in (_get_undo_file(self.folder),):
                        try:
                            _tmp = _undo_path + ".tmp"
                            with open(_tmp, "w", encoding="utf-8") as f:
                                json.dump({"folder": self.folder, "renames": undo_map}, f, indent=2, ensure_ascii=False)
                                f.flush()
                                try:
                                    os.fsync(f.fileno())
                                except Exception:
                                    pass
                            os.replace(_tmp, _undo_path)
                        except Exception:
                            pass

                except PermissionError:
                    error_count += 1
                    error_files.append(f"{fname} (soubor je otevřen jiným programem)")
                    logger.error("Soubor %s je uzamčen jiným procesem", fname)
                except Exception as e:
                    error_count += 1
                    error_files.append(f"{fname} ({e})")
                    logger.error("Chyba přejmenování %s: %s", fname, e)
            else:
                # src == dst → již správně pojmenováno, považujeme za úspěch pro Excel
                success = True
                # nezvyšujeme renamed_count (žádný rename), ale pro Excel ano

            if success:
                successful_fnames.add(fname)
                successful_all.append((safe_num, name_val, fmt_val))
                if str(conf).startswith("Ověřeno"):
                    successful_verified.append((safe_num, name_val, fmt_val))

        # Excel až po úspěšném rename — why: původně i neúspěšné byly v Excelu → nesoulad FS
        verified_export_items = successful_verified if successful_verified else successful_all

        # Automatický zápis do Excelu
        excel_path = os.path.join(self.folder, "seznam_vykresu.xlsx")
        excel_msg = ""
        if verified_export_items:
            try:
                export_to_excel(excel_path, verified_export_items, self.paper_format)
                excel_msg = f"\n\n📊 Do Excelu bylo zapsáno {len(verified_export_items)} výkresů:\n{excel_path}"
            except Exception as e:
                logger.error("Chyba automatického exportu do Excelu: %s", e)
                excel_msg = f"\n\n⚠️ Varování: Export do Excelu selhal: {e}"

        # Závěrečná zpráva s počtem úspěchů i chyb
        result_msg = f"✅ Úspěšně přejmenováno: {renamed_count} souborů"
        if error_count > 0:
            result_msg += f"\n❌ Nepodařilo se přejmenovat: {error_count} souborů"
            # Zobrazit max 10 chybových souborů
            shown_errors = error_files[:10]
            result_msg += "\n" + "\n".join(f"  • {ef}" for ef in shown_errors)
            if len(error_files) > 10:
                result_msg += f"\n  ... a dalších {len(error_files) - 10}"

        # Po úspěšném přejmenování vyčistit nebo aktualizovat checkpoint — why: původně mazal i neúspěšné
        if successful_fnames:
            chk = get_folder_checkpoint(self.folder)
            if chk and "results" in chk:
                for fname in list(successful_fnames):
                    if fname in chk["results"]:
                        del chk["results"][fname]
                if not chk["results"]:
                    remove_folder_checkpoint(self.folder)
                else:
                    save_folder_checkpoint(self.folder, chk)
            else:
                remove_folder_checkpoint(self.folder)

        if undo_map:
            result_msg += f"\n\n🔄 Záloha pro vrácení změn uložena."

        result_msg += excel_msg

        if error_count > 0:
            messagebox.showwarning("Dokončeno s chybami", result_msg)
        else:
            messagebox.showinfo("Hotovo", result_msg)
        self.on_exit()

    def on_exit(self):
        """Bezpečně ukončí vlákna a zavře aplikaci."""
        try:
            if hasattr(self, '_pause_requested'):
                self._pause_requested.set()
        except Exception:
            pass
        try:
            if hasattr(self, 'preview_executor'):
                self.preview_executor.shutdown(wait=False, cancel_futures=True)
        except Exception:
            pass
        try:
            self.root.quit()
        except Exception:
            pass
        try:
            self.root.destroy()
        except Exception:
            pass
